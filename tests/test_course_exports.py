"""Independent ZIP/XML/relationship and python-docx/python-pptx interoperability validation."""
import os, pathlib, posixpath, unittest, zipfile, xml.etree.ElementTree as ET
from docx import Document
from pptx import Presentation

ROOT=pathlib.Path(os.environ.get('QCU_COURSE_EXPORTS','.work/course-validation/exports'))
class Exports(unittest.TestCase):
    def test_office_archives(self):
        paths=list(ROOT.glob('*.docx'))+list(ROOT.glob('*.pptx'))
        self.assertEqual(len(paths),7)
        for path in paths:
            with self.subTest(path=path.name),zipfile.ZipFile(path) as z:
                self.assertIsNone(z.testzip());names=set(z.namelist())
                for name in names:
                    root=ET.fromstring(z.read(name))
                    if name.endswith('.rels'):
                        base=posixpath.dirname(posixpath.dirname(name))
                        for rel in root:
                            self.assertNotEqual(rel.attrib.get('TargetMode'),'External')
                            target=posixpath.normpath(posixpath.join(base,rel.attrib['Target']))
                            self.assertIn(target,names,(name,target))
                if path.suffix=='.docx':
                    d=Document(path);text='\n'.join(p.text for p in d.paragraphs)
                    self.assertIn('教师审核草案',text);self.assertIn('概念理解',text);self.assertGreater(len(d.paragraphs),10)
                else:
                    p=Presentation(path);self.assertGreater(len(p.slides),5);self.assertEqual(p.slide_width,12192000)
                    texts='\n'.join(s.text for slide in p.slides for s in slide.shapes if s.has_text_frame)
                    self.assertIn('Quadratic',texts);self.assertIn('二次函数',texts)
    def test_offline_html(self):
        paths=list(ROOT.glob('*.html'));self.assertEqual(len(paths),8)
        for path in paths:
            html=path.read_text();self.assertIn("connect-src 'none'",html);self.assertNotIn('cdn.',html)
            self.assertIn('Teacher review required',html)

if __name__=='__main__':unittest.main()
